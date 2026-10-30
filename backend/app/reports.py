"""Reports: tabular views of jobs, candidates, closures and performance.

Every report is shown 100 rows per page at most, and can be downloaded as an Excel
file containing *all* rows that match the filters. For your team only (Super Admin and
recruiters); partner and client logins cannot open reports.

Deactivated jobs (e.g. test data) are left out, as on the dashboard.
"""
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Callable

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.db_dependencies import get_db
from app.dependencies import require_permission
from app.models import (
    CandidateStatusHistory,
    CandidateSubmission,
    Client,
    Job,
    JobVendorAssignment,
    Role,
    User,
    Vendor,
)
from app.partner_access import is_client_user, is_partner_user

router = APIRouter(prefix="/reports", tags=["Reports"])

PAGE_SIZE_MAX = 100
CLOSED = {"Selected", "Offer Accepted", "Joined"}
IN_PROCESS = {"Submitted", "Screening", "Shortlisted", "Sent to Client", "Interview Scheduled"}
SHORTLISTED_OR_LATER = {"Shortlisted", "Sent to Client", "Interview Scheduled", "Selected", "Offer Accepted", "Joined"}


@dataclass
class Column:
    key: str
    label: str
    type: str = "text"   # text | number | date | money | percent
    total: bool = False  # summed in the Excel totals row
    # percent columns are computed in Excel as numerator / denominator (keys of other columns)
    ratio: tuple[str, str] | None = None


@dataclass
class Report:
    key: str
    title: str
    description: str
    columns: list[Column]
    build: Callable[["Ctx"], list[dict]]
    summary: bool = False  # adds a totals row in Excel


def _utc(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value if value.tzinfo else value.replace(tzinfo=timezone.utc)


class Ctx:
    """Everything a report needs, loaded once per request."""

    def __init__(self, db: Session, user: User, start: datetime | None, end: datetime | None, client_id: int | None):
        self.db, self.user = db, user
        self.start, self.end = _utc(start), _utc(end)
        org = user.organization_id
        job_query = (
            select(Job)
            .options(selectinload(Job.client), selectinload(Job.recruiter))
            .where(Job.organization_id == org, Job.is_active.is_(True))
        )
        if client_id is not None:
            job_query = job_query.where(Job.client_id == client_id)
        self.jobs = db.scalars(job_query.order_by(Job.id.desc())).all()
        self.job_by_id = {job.id: job for job in self.jobs}
        self.subs = db.scalars(
            select(CandidateSubmission)
            .options(selectinload(CandidateSubmission.candidate), selectinload(CandidateSubmission.vendor))
            .where(CandidateSubmission.job_id.in_(list(self.job_by_id) or [-1]))
            .order_by(CandidateSubmission.created_at.desc(), CandidateSubmission.id.desc())
        ).all()
        self.reached: dict[int, set[str]] = {}
        self.first_closed: dict[int, datetime] = {}
        for submission_id, status, changed_at in db.execute(
            select(CandidateStatusHistory.submission_id, CandidateStatusHistory.new_status, CandidateStatusHistory.changed_at)
            .where(CandidateStatusHistory.submission_id.in_([item.id for item in self.subs] or [-1]))
        ).all():
            self.reached.setdefault(submission_id, set()).add(status)
            if status in CLOSED:
                changed_at = _utc(changed_at)
                if submission_id not in self.first_closed or changed_at < self.first_closed[submission_id]:
                    self.first_closed[submission_id] = changed_at
        self.assignments = db.execute(
            select(JobVendorAssignment.job_id, JobVendorAssignment.vendor_id).where(
                JobVendorAssignment.job_id.in_(list(self.job_by_id) or [-1]), JobVendorAssignment.is_active.is_(True)
            )
        ).all()
        self.vendors = db.scalars(select(Vendor).where(Vendor.organization_id == org).order_by(Vendor.name)).all()
        clients_query = select(Client).where(Client.organization_id == org)
        if client_id is not None:
            clients_query = clients_query.where(Client.id == client_id)
        self.clients = db.scalars(clients_query.order_by(Client.company_name)).all()
        self.recruiters = db.scalars(
            select(User).join(User.roles).where(User.organization_id == org, Role.name == "Recruiter").order_by(User.id)
        ).unique().all()

    def in_period(self, moment: datetime | None) -> bool:
        moment = _utc(moment)
        if moment is None:
            return False
        return (self.start is None or moment >= self.start) and (self.end is None or moment < self.end)

    def period_subs(self) -> list[CandidateSubmission]:
        return [item for item in self.subs if self.in_period(item.created_at)]

    def job_fields(self, job_id: int) -> dict:
        job = self.job_by_id[job_id]
        return {
            "position": job.title,
            "client": job.client.company_name if job.client else None,
            "recruiter": (job.recruiter.full_name or job.recruiter.email) if job.recruiter else None,
        }

    def ever(self, submission: CandidateSubmission, statuses: set[str]) -> bool:
        return bool((self.reached.get(submission.id, set()) | {submission.status}) & statuses)


def _counts(ctx: Ctx, subs: list[CandidateSubmission]) -> dict:
    return {
        "submissions": len(subs),
        "shortlisted": sum(1 for item in subs if ctx.ever(item, SHORTLISTED_OR_LATER)),
        "sent_to_client": sum(1 for item in subs if ctx.ever(item, {"Sent to Client"})),
        "interviews": sum(1 for item in subs if ctx.ever(item, {"Interview Scheduled"})),
        "closures": sum(1 for item in subs if item.status in CLOSED),
        "joined": sum(1 for item in subs if item.status == "Joined"),
    }


def _rate(numerator: int, denominator: int) -> float:
    return round(numerator / denominator, 4) if denominator else 0.0


# ---------------------------------------------------------------------------
# Report builders
# ---------------------------------------------------------------------------

def build_submissions(ctx: Ctx) -> list[dict]:
    rows = []
    for item in ctx.period_subs():
        rows.append({
            "candidate": item.candidate.full_name,
            "email": item.candidate.email,
            "phone": item.candidate.phone,
            **ctx.job_fields(item.job_id),
            "partner": item.vendor.name if item.vendor else "Internal",
            "status": item.status,
            "experience": item.candidate.total_experience,
            "employer": item.candidate.current_employer,
            "location": item.candidate.current_location,
            "submitted_on": _utc(item.created_at),
            "updated_on": _utc(item.updated_at),
        })
    return rows


def build_closures(ctx: Ctx) -> list[dict]:
    rows = []
    for item in ctx.subs:
        closed_on = ctx.first_closed.get(item.id)
        if item.status in CLOSED and ctx.in_period(closed_on):
            rows.append({
                "candidate": item.candidate.full_name,
                **ctx.job_fields(item.job_id),
                "partner": item.vendor.name if item.vendor else "Internal",
                "stage": item.status,
                "closed_on": closed_on,
            })
    rows.sort(key=lambda row: row["closed_on"], reverse=True)
    return rows


def build_jobs(ctx: Ctx) -> list[dict]:
    rows = []
    for job in ctx.jobs:
        if not ctx.in_period(job.created_at):
            continue
        subs = [item for item in ctx.subs if item.job_id == job.id]
        counts = _counts(ctx, subs)
        rows.append({
            "job_id": job.id,
            "position": job.title,
            "client": job.client.company_name if job.client else None,
            "recruiter": (job.recruiter.full_name or job.recruiter.email) if job.recruiter else None,
            "status": job.status.capitalize(),
            "openings": job.openings,
            "min_exp": job.min_experience,
            "max_exp": job.max_experience,
            "qualification": job.qualification,
            "ctc_min": job.budget_min,
            "ctc_max": job.budget_max,
            "partners": sum(1 for job_id, _vendor in ctx.assignments if job_id == job.id),
            "submissions": counts["submissions"],
            "sent_to_client": counts["sent_to_client"],
            "closures": counts["closures"],
            "created_on": _utc(job.created_at),
        })
    return rows


def build_partners(ctx: Ctx) -> list[dict]:
    period = ctx.period_subs()
    rows = []
    for vendor in ctx.vendors:
        counts = _counts(ctx, [item for item in period if item.vendor_id == vendor.id])
        rows.append({
            "partner": vendor.name,
            "status": "Active" if vendor.is_active else "Inactive",
            "positions": sum(1 for _job, vendor_id in ctx.assignments if vendor_id == vendor.id),
            **{key: counts[key] for key in ("submissions", "shortlisted", "sent_to_client", "interviews", "closures")},
            "closure_rate": _rate(counts["closures"], counts["submissions"]),
        })
    rows.sort(key=lambda row: (-row["closures"], -row["submissions"], row["partner"].lower()))
    return rows


def build_recruiters(ctx: Ctx) -> list[dict]:
    period = ctx.period_subs()
    rows = []
    for recruiter in ctx.recruiters:
        own = {job.id for job in ctx.jobs if job.recruiter_id == recruiter.id}
        counts = _counts(ctx, [item for item in period if item.job_id in own])
        rows.append({
            "recruiter": recruiter.full_name or recruiter.email,
            "email": recruiter.email,
            "status": "Active" if recruiter.is_active else "Inactive",
            "positions": len(own),
            "open_positions": sum(1 for job in ctx.jobs if job.id in own and job.status == "open"),
            **{key: counts[key] for key in ("submissions", "sent_to_client", "closures", "joined")},
            "closure_rate": _rate(counts["closures"], counts["submissions"]),
        })
    rows.sort(key=lambda row: (-row["closures"], -row["submissions"], row["recruiter"].lower()))
    return rows


def build_clients(ctx: Ctx) -> list[dict]:
    period = ctx.period_subs()
    rows = []
    for client in ctx.clients:
        client_jobs = [job for job in ctx.jobs if job.client_id == client.id]
        ids = {job.id for job in client_jobs}
        counts = _counts(ctx, [item for item in period if item.job_id in ids])
        rows.append({
            "client": client.company_name,
            "status": "Active" if client.is_active else "Inactive",
            "positions": len(client_jobs),
            "open_positions": sum(1 for job in client_jobs if job.status == "open"),
            **{key: counts[key] for key in ("submissions", "sent_to_client", "interviews", "closures")},
            "closure_rate": _rate(counts["closures"], counts["submissions"]),
        })
    rows.sort(key=lambda row: (-row["closures"], -row["submissions"], row["client"].lower()))
    return rows


C = Column
REPORTS: dict[str, Report] = {
    report.key: report
    for report in [
        Report("submissions", "Candidate submissions", "Every candidate submitted in the period, with their position, client, partner, recruiter and current status.", [
            C("candidate", "Candidate"), C("email", "Email"), C("phone", "Phone"), C("position", "Position"), C("client", "Client"),
            C("partner", "Partner"), C("recruiter", "Recruiter"), C("status", "Status"), C("experience", "Total exp. (yrs)", "number"),
            C("employer", "Current employer"), C("location", "Location"), C("submitted_on", "Submitted on", "date"), C("updated_on", "Last updated", "date"),
        ], build_submissions),
        Report("closures", "Closures", "Candidates who reached Selected, Offer Accepted or Joined in the period (dated by when they first got there).", [
            C("candidate", "Candidate"), C("position", "Position"), C("client", "Client"), C("partner", "Partner"),
            C("recruiter", "Recruiter"), C("stage", "Stage"), C("closed_on", "Closed on", "date"),
        ], build_closures),
        Report("jobs", "Positions", "Every position created in the period, with its owner, budget and how far its candidates have got (all time).", [
            C("job_id", "Job ID", "number"), C("position", "Position"), C("client", "Client"), C("recruiter", "Recruiter"), C("status", "Status"),
            C("openings", "Openings", "number", True), C("min_exp", "Min exp. (yrs)", "number"), C("max_exp", "Max exp. (yrs)", "number"),
            C("qualification", "Qualification"), C("ctc_min", "Offered CTC min (₹/yr)", "money"), C("ctc_max", "Offered CTC max (₹/yr)", "money"),
            C("partners", "Partners", "number"), C("submissions", "Submissions", "number", True), C("sent_to_client", "Sent to client", "number", True),
            C("closures", "Closures", "number", True), C("created_on", "Created on", "date"),
        ], build_jobs, summary=True),
        Report("partners", "Partner performance", "For candidates submitted in the period: how far each partner's candidates got.", [
            C("partner", "Partner"), C("status", "Status"), C("positions", "Positions assigned", "number", True),
            C("submissions", "Submissions", "number", True), C("shortlisted", "Shortlisted or later", "number", True),
            C("sent_to_client", "Sent to client", "number", True), C("interviews", "Interviews", "number", True),
            C("closures", "Closures", "number", True), C("closure_rate", "Closure rate", "percent", ratio=("closures", "submissions")),
        ], build_partners, summary=True),
        Report("recruiters", "Recruiter performance", "For candidates submitted in the period on each recruiter's positions.", [
            C("recruiter", "Recruiter"), C("email", "Email"), C("status", "Status"), C("positions", "Positions", "number", True),
            C("open_positions", "Open positions", "number", True), C("submissions", "Submissions", "number", True),
            C("sent_to_client", "Sent to client", "number", True), C("closures", "Closures", "number", True), C("joined", "Joined", "number", True),
            C("closure_rate", "Closure rate", "percent", ratio=("closures", "submissions")),
        ], build_recruiters, summary=True),
        Report("clients", "Client summary", "For each client: positions, and for candidates submitted in the period, how far they got.", [
            C("client", "Client"), C("status", "Status"), C("positions", "Positions", "number", True), C("open_positions", "Open positions", "number", True),
            C("submissions", "Submissions", "number", True), C("sent_to_client", "Sent to client", "number", True),
            C("interviews", "Interviews", "number", True), C("closures", "Closures", "number", True),
            C("closure_rate", "Closure rate", "percent", ratio=("closures", "submissions")),
        ], build_clients, summary=True),
    ]
}


# ---------------------------------------------------------------------------
# Endpoints
# ---------------------------------------------------------------------------

def _staff(current_user: User = Depends(require_permission("candidates.view"))) -> User:
    if is_partner_user(current_user) or is_client_user(current_user):
        raise HTTPException(status_code=403, detail="You do not have permission to perform this action")
    return current_user


def _rows(report: Report, ctx: Ctx, q: str | None) -> list[dict]:
    rows = report.build(ctx)
    needle = (q or "").strip().lower()
    if needle:
        text_keys = [column.key for column in report.columns if column.type == "text"]
        rows = [row for row in rows if any(needle in str(row.get(key) or "").lower() for key in text_keys)]
    return rows


def _report_or_404(key: str) -> Report:
    report = REPORTS.get(key)
    if report is None:
        raise HTTPException(status_code=404, detail="Report not found")
    return report


@router.get("")
def list_reports(_user: User = Depends(_staff)):
    return [{"key": report.key, "title": report.title, "description": report.description} for report in REPORTS.values()]


@router.get("/{key}")
def get_report(
    key: str,
    start: datetime | None = None,
    end: datetime | None = None,
    client_id: int | None = None,
    q: str | None = Query(default=None, max_length=200),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=PAGE_SIZE_MAX, ge=1, le=PAGE_SIZE_MAX),
    db: Session = Depends(get_db),
    current_user: User = Depends(_staff),
):
    report = _report_or_404(key)
    rows = _rows(report, Ctx(db, current_user, start, end, client_id), q)
    total = len(rows)
    first = (page - 1) * page_size
    return {
        "key": report.key,
        "title": report.title,
        "description": report.description,
        "columns": [{"key": c.key, "label": c.label, "type": c.type} for c in report.columns],
        "rows": rows[first:first + page_size],
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": max(1, -(-total // page_size)),
    }


@router.get("/{key}/export")
def export_report(
    key: str,
    start: datetime | None = None,
    end: datetime | None = None,
    client_id: int | None = None,
    q: str | None = Query(default=None, max_length=200),
    tz_offset: int = Query(default=0, ge=-840, le=840, description="Minutes east of UTC, for dates in the file"),
    db: Session = Depends(get_db),
    current_user: User = Depends(_staff),
):
    """All rows that match the filters, as an Excel workbook."""
    report = _report_or_404(key)
    ctx = Ctx(db, current_user, start, end, client_id)
    rows = _rows(report, ctx, q)
    client_name = next((c.company_name for c in ctx.clients if c.id == client_id), None) if client_id else None
    content = _workbook(report, rows, ctx, client_name, q, tz_offset)
    local_now = datetime.now(timezone.utc) + timedelta(minutes=tz_offset)
    filename = f"{report.key}-report-{local_now:%Y-%m-%d}.xlsx"
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _workbook(report: Report, rows: list[dict], ctx: Ctx, client_name: str | None, q: str | None, tz_offset: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    shift = timedelta(minutes=tz_offset)
    local = lambda moment: (moment + shift).replace(tzinfo=None) if moment else None
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="173153")
    total_fill = PatternFill("solid", fgColor="EEF2F7")
    thin = Side(style="thin", color="D5DDE7")

    wb = Workbook()
    ws = wb.active
    ws.title = report.title[:31]
    columns = [Column("sno", "S.No", "number")] + report.columns
    letter = {column.key: get_column_letter(index + 1) for index, column in enumerate(columns)}

    def period_text() -> str:
        fmt = lambda moment: (moment + shift).strftime("%d %b %Y")
        if ctx.start and ctx.end:
            return f"{fmt(ctx.start)} – {fmt(ctx.end - timedelta(seconds=1))}"
        if ctx.start:
            return f"from {fmt(ctx.start)}"
        if ctx.end:
            return f"until {fmt(ctx.end - timedelta(seconds=1))}"
        return "All time"

    ws["A1"] = report.title
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="173153")
    ws["A2"] = f"Period: {period_text()}   ·   Client: {client_name or 'All clients'}   ·   Search: {q.strip() if q and q.strip() else '—'}"
    generated = (datetime.now(timezone.utc) + shift).strftime("%d %b %Y, %I:%M %p")
    ws["A3"] = f"Generated {generated} by {ctx.user.full_name or ctx.user.email}   ·   {len(rows)} row(s)"
    for cell in ("A2", "A3"):
        ws[cell].font = Font(name="Arial", size=9, color="5B6B80")

    header_row = 5
    for index, column in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=index, value=column.label)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="center" if column.key == "sno" else None)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[header_row].height = 30

    formats = {"date": "dd mmm yyyy", "money": '"₹"#,##0', "percent": "0.0%", "number": "#,##0.##"}
    first_data = header_row + 1
    for offset, row in enumerate(rows):
        r = first_data + offset
        for index, column in enumerate(columns, start=1):
            if column.key == "sno":
                value = offset + 1
            elif column.ratio:
                num, den = (f"{letter[key]}{r}" for key in column.ratio)
                value = f"=IF({den}=0,0,{num}/{den})"  # live formula: recalculates if the numbers are edited
            else:
                value = row.get(column.key)
                if column.type == "date":
                    value = local(value)
            cell = ws.cell(row=r, column=index, value=value)
            cell.font = font
            if column.key == "sno":
                cell.alignment = Alignment(horizontal="center")
            if column.type in formats:
                cell.number_format = formats[column.type]
            cell.border = Border(bottom=thin)
    last_data = first_data + len(rows) - 1

    if report.summary and rows:
        total_row = last_data + 1
        for index, column in enumerate(columns, start=1):
            cell = ws.cell(row=total_row, column=index)
            if index == 2:
                cell.value = "Total"
            elif column.total:
                # SUBTOTAL(109, …) ignores rows hidden by the filter arrows
                cell.value = f"=SUBTOTAL(109,{letter[column.key]}{first_data}:{letter[column.key]}{last_data})"
                cell.number_format = formats["number"]
            elif column.ratio:
                num, den = (f"{letter[key]}{total_row}" for key in column.ratio)
                cell.value = f"=IF({den}=0,0,{num}/{den})"
                cell.number_format = formats["percent"]
            cell.font, cell.fill = bold, total_fill

    ws.freeze_panes = ws.cell(row=first_data, column=2)
    if rows:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{last_data}"
    for index, column in enumerate(columns, start=1):
        sample = [len(column.label)] + [len(str(row.get(column.key) or "")) for row in rows[:500]]
        width = 8 if column.key == "sno" else 12 if column.type in ("date", "number", "percent") else min(max(sample) + 2, 45)
        ws.column_dimensions[get_column_letter(index)].width = max(width, min(len(column.label) + 2, 24))

    notes = wb.create_sheet("About this report")
    lines = [
        (report.title, Font(name="Arial", size=13, bold=True, color="173153")),
        (report.description, font),
        ("", font),
        ("Definitions", bold),
        ("Closure: a candidate whose current stage is Selected, Offer Accepted or Joined.", font),
        ("Sent to client / Interviews / Shortlisted or later: the candidate reached that stage at some point.", font),
        ("Closure rate: closures ÷ submissions (a formula, so it updates if you edit the numbers).", font),
        ("Deactivated jobs (e.g. test data) are not included.", font),
        ("", font),
        (f"Period: {period_text()}", font),
        (f"Client: {client_name or 'All clients'}", font),
        (f"Generated: {generated}", font),
    ]
    for index, (text, style) in enumerate(lines, start=1):
        notes.cell(row=index, column=1, value=text).font = style
    notes.column_dimensions["A"].width = 110

    wb.calculation.fullCalcOnLoad = True  # Excel computes the formulas when the file is opened
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
