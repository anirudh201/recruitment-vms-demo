"""Read a PDF/DOCX resume and suggest values for the candidate form.

Rules are deliberately conservative: a field is only filled when the resume states it
clearly. Anything uncertain is left empty for the person submitting to type in.
"""
import re
from io import BytesIO

EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")
PHONE_CANDIDATE = re.compile(r"\+?\d[\d\s().-]{8,17}\d")
YEAR_RANGE = re.compile(r"^\s*(?:19|20)\d{2}(?:\s*[-–to/]+\s*(?:19|20)\d{2})*\s*$")

SECTION_HEADINGS = re.compile(
    r"^(?:professional\s+|work\s+|employment\s+|career\s+)?(?:experience|history|summary|profile|objective|education|"
    r"academic(?:s|\s+details|\s+qualifications?)?|qualifications?|projects?|certifications?|achievements?|awards?|"
    r"personal\s+(?:details|information|profile)|languages?(?:\s+known)?|declaration|hobbies|interests|"
    r"(?:technical\s+|key\s+|core\s+|primary\s+|it\s+)?skills?(?:\s+set|\s+summary)?|strengths|responsibilities|"
    r"professional\s+experience|work\s+experience|career\s+objective|contact(?:\s+details)?)\s*:?$",
    re.I,
)
SKILL_HEADING = re.compile(
    r"^(?:technical\s+|key\s+|core\s+|primary\s+|it\s+|professional\s+)?skills?(?:\s+set|\s+summary)?\s*:?\s*(.*)$", re.I
)
NAME_STOPWORDS = {
    "resume", "curriculum", "vitae", "cv", "profile", "summary", "objective", "contact", "email", "phone",
    "mobile", "address", "experience", "education", "skills", "career", "professional", "personal", "details",
}
COMPANY_HINT = re.compile(
    r"\b(?:ltd|limited|pvt|private|inc|llp|llc|corp|corporation|technologies|technology|solutions|systems|services|"
    r"consulting|consultancy|infotech|software|labs|global|india|group|bank|tech)\b\.?",
    re.I,
)
CURRENT_MARK = re.compile(r"\b(?:present|till\s+date|to\s+date|till\s+now|current(?:ly)?|ongoing)\b", re.I)
MONTH = r"(?:jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?"
DATE_TOKEN = re.compile(rf"(?:{MONTH}\s*)?(?:19|20)\d{{2}}|{MONTH}", re.I)

CITIES = [
    "Hyderabad", "Secunderabad", "Bengaluru", "Bangalore", "Chennai", "Pune", "Mumbai", "Navi Mumbai", "Thane",
    "New Delhi", "Delhi", "Noida", "Gurugram", "Gurgaon", "Kolkata", "Ahmedabad", "Kochi", "Cochin", "Coimbatore",
    "Visakhapatnam", "Vizag", "Vijayawada", "Jaipur", "Indore", "Chandigarh", "Trivandrum", "Thiruvananthapuram",
    "Mysuru", "Mysore", "Nagpur", "Lucknow", "Bhubaneswar", "Warangal", "Madurai", "Mangalore", "Vadodara", "Surat",
]
# highest first
DEGREES = [
    ("Ph.D", re.compile(r"\bph\.?\s?d\b|\bdoctorate\b", re.I)),
    ("M.Tech", re.compile(r"\bm\.?\s?tech\b", re.I)),
    ("M.E.", re.compile(r"\bM\.E\.?(?=[\s,(]|$)|\bmaster\s+of\s+engineering\b")),
    ("MBA", re.compile(r"\bmba\b|\bpgdm\b|\bmaster\s+of\s+business", re.I)),
    ("MCA", re.compile(r"\bmca\b|\bmaster\s+of\s+computer\s+applications?\b", re.I)),
    ("M.Sc", re.compile(r"\bm\.?\s?sc\b|\bmaster\s+of\s+science\b", re.I)),
    ("M.Com", re.compile(r"\bm\.?\s?com\b|\bmaster\s+of\s+commerce\b", re.I)),
    ("B.Tech", re.compile(r"\bb\.?\s?tech\b|\bbachelor\s+of\s+technology\b", re.I)),
    ("B.E.", re.compile(r"\bB\.E\.?(?=[\s,(]|$)|\bbachelor\s+of\s+engineering\b")),
    ("BCA", re.compile(r"\bbca\b|\bbachelor\s+of\s+computer\s+applications?\b", re.I)),
    ("B.Sc", re.compile(r"\bb\.?\s?sc\b|\bbachelor\s+of\s+science\b", re.I)),
    ("B.Com", re.compile(r"\bb\.?\s?com\b|\bbachelor\s+of\s+commerce\b", re.I)),
    ("BBA", re.compile(r"\bbba\b|\bbachelor\s+of\s+business", re.I)),
    ("Diploma", re.compile(r"\bdiploma\b", re.I)),
]
SPECIALIZATION = re.compile(
    r"\b(computer\s+science(?:\s+(?:and|&)\s+engineering)?|information\s+technology|electronics(?:\s+(?:and|&)\s+"
    r"communication(?:s)?)?(?:\s+engineering)?|electrical(?:\s+(?:and|&)\s+electronics)?(?:\s+engineering)?|"
    r"mechanical(?:\s+engineering)?|civil(?:\s+engineering)?|chemical(?:\s+engineering)?|finance|marketing|"
    r"human\s+resources|data\s+science|artificial\s+intelligence|mathematics|physics|statistics)\b",
    re.I,
)
KNOWN_SKILLS = [
    "Java", "Python", "JavaScript", "TypeScript", "C#", ".NET", "C++", "Go", "Kotlin", "Swift", "PHP", "Ruby", "Scala",
    "SQL", "PL/SQL", "PostgreSQL", "MySQL", "Oracle", "MongoDB", "Redis", "SQL Server",
    "React", "Angular", "Vue.js", "Node.js", "Next.js", "Django", "Flask", "FastAPI", "Spring Boot", "Spring", "Hibernate",
    "HTML", "CSS", "REST", "GraphQL", "Microservices",
    "AWS", "Azure", "GCP", "Docker", "Kubernetes", "Terraform", "Jenkins", "Git", "Linux", "DevOps", "CI/CD",
    "Kafka", "Spark", "Hadoop", "Airflow", "Snowflake", "Databricks", "Power BI", "Tableau", "Excel",
    "Machine Learning", "Deep Learning", "Data Science", "NLP", "TensorFlow", "PyTorch",
    "Selenium", "JUnit", "Cypress", "Jira", "Agile", "Scrum",
    "SAP", "SAP FICO", "SAP MM", "SAP SD", "SAP EWM", "SAP ABAP", "S/4HANA", "Oracle PLM", "Oracle Cloud", "Salesforce",
    "ServiceNow", "Workday",
]


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def _clean(text: str) -> list[str]:
    text = text.replace("\r", "\n")
    text = re.sub(r"[\x0b\x0c]", "\n", text)
    text = re.sub(r"[\x00-\x08\x0e-\x1f\x7f\uf000-\uf8ff]", "", text)
    lines = [re.sub(r"\s+", " ", line).strip(" \t•*·▪●○◦-–:|") for line in text.splitlines()]
    return [line for line in lines if line]


def extract_text_variants(filename: str, content: bytes) -> list[list[str]]:
    """Lines of the document, possibly in two readings (PDF layout vs reading order)."""
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(content))
        variants = []
        for kwargs in ({}, {"extraction_mode": "layout"}):
            try:
                text = "\n".join((page.extract_text(**kwargs) or "") for page in reader.pages)
            except Exception:
                continue
            if text.strip():
                variants.append(_clean(text))
        return variants
    from docx import Document

    document = Document(BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:  # many resumes are laid out in tables
        for row in table.rows:
            seen = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in seen:   # merged cells repeat their text
                    seen.append(value)
            if len(seen) == 2 and len(seen[0]) <= 30 and "\n" not in seen[0]:
                parts.append(f"{seen[0].rstrip(':')}: {seen[1]}")   # "Total Experience | 6 years" -> one line
            else:
                parts.extend(seen)
    for section in document.sections:  # name/contact often sit in the header
        parts = [p.text for p in section.header.paragraphs] + parts
    return [_clean("\n".join(parts))]


# ---------------------------------------------------------------------------
# Field detection
# ---------------------------------------------------------------------------

def _email(lines: list[str]) -> str | None:
    for line in lines:
        match = EMAIL.search(line)
        if match:
            return match.group(0).lower()
    return None


def _phone(lines: list[str]) -> str | None:
    best: tuple[int, int, str] | None = None
    order = 0
    for line in lines[:40]:
        for match in PHONE_CANDIDATE.finditer(line):
            raw = match.group(0).strip()
            if YEAR_RANGE.match(raw):
                continue
            digits = re.sub(r"\D", "", raw)
            if not 10 <= len(digits) <= 13:
                continue
            local = digits[-10:]
            if len(digits) == 10 and local[0] not in "6789" and not raw.startswith("+"):
                continue  # not an Indian mobile and no country code: likely an ID/date
            score = 0
            if local[0] in "6789" and not digits.startswith("0"):
                score += 2                                   # looks like a mobile number
            if digits.startswith("0"):
                score -= 2                                   # STD code: a landline
            before = line[max(0, match.start() - 18):match.start()].lower()
            if re.search(r"mob|cell|whats\s?app|phone|contact\s*(?:no|number)", before):
                score += 1
            if re.search(r"office|landline|fax|res\b", line[match.end():match.end() + 12].lower() + before):
                score -= 1
            order += 1
            candidate = (score, -order, re.sub(r"\s{2,}", " ", raw))
            if best is None or candidate > best:
                best = candidate
    return best[2] if best else None


def _looks_like_name(line: str) -> bool:
    if any(ch.isdigit() for ch in line) or "@" in line or len(line) > 40:
        return False
    words = line.replace(".", " ").split()
    if not 2 <= len(words) <= 4:
        return False
    if any(word.lower().strip(",") in NAME_STOPWORDS for word in words):
        return False
    return all(re.fullmatch(r"[A-Za-z][A-Za-z'.-]*", word) for word in words)


def _tidy_name(name: str) -> str:
    return " ".join(w.capitalize() if w.isupper() or w.islower() else w for w in name.split())


def _name(lines: list[str]) -> str | None:
    for line in lines[:30]:
        match = re.match(r"^(?:full\s+)?name\s*[:\-]\s*(.+)$", line, re.I)
        if match and _looks_like_name(match.group(1).strip()):
            return _tidy_name(match.group(1).strip())
    for line in lines[:8]:
        candidate = re.split(r"\s+[|,]\s+", line)[0].strip()
        if _looks_like_name(candidate):
            return _tidy_name(candidate)
    return None


def _years(text: str) -> float | None:
    match = re.search(r"(\d{1,2}(?:\.\d{1,2})?)\s*\+?\s*(?:years?|yrs?)(?:\s*(?:and\s*)?(\d{1,2})\s*(?:months?|mos?))?", text, re.I)
    if not match:
        return None
    value = float(match.group(1)) + (int(match.group(2)) / 12 if match.group(2) else 0)
    return round(value, 1) if 0 <= value <= 50 else None


def _experience(lines: list[str]) -> tuple[float | None, float | None]:
    total = relevant = None
    for line in lines:
        if not re.search(r"\bexp(?:erience)?\b", line, re.I):
            continue
        value = _years(line)
        if value is None:
            continue
        if re.search(r"\brelevant\b", line, re.I):
            relevant = relevant if relevant is not None else value
        elif total is None:
            total = value
    return total, relevant


def _employer(lines: list[str]) -> str | None:
    patterns = [
        re.compile(r"current(?:ly)?\s+(?:company|employer|organi[sz]ation)\s*[:\-]\s*(.+)", re.I),
        re.compile(r"current(?:ly)?\s+(?:working|employed)\s+(?:at|with|in)\s+(.+)", re.I),
        re.compile(r"\bworking\s+(?:at|with)\s+([A-Z][\w&.,'\- ]{1,60})", re.I),
    ]
    for line in lines:
        for pattern in patterns:
            match = pattern.search(line)
            if match:
                company = re.split(r"\s+(?:as|since|from)\s+|[,(]|\s+[-–|]\s+|\.\s", match.group(1))[0].strip(" .")
                if 1 <= len(company.split()) <= 8:
                    return company
    # experience entries marked "Present" / "Till date"
    for index, line in enumerate(lines):
        if not CURRENT_MARK.search(line) or not DATE_TOKEN.search(line):
            continue
        for candidate_line in (line, lines[index - 1] if index else ""):
            without_dates = re.sub(rf"(?:{MONTH}\s*)?(?:19|20)\d{{2}}|{MONTH}|{CURRENT_MARK.pattern}", " ", candidate_line, flags=re.I)
            parts = [part.strip(" ,.-–()") for part in re.split(r"\s+[|–-]\s+|\||,|\(|\)|\s{3,}", without_dates)]
            parts = [part for part in parts if part and re.search(r"[A-Za-z]{2,}", part)]
            named = [part for part in parts if COMPANY_HINT.search(part)]
            if named:
                return named[0]
    return None


def _location(lines: list[str]) -> str | None:
    for line in lines[:40]:
        match = re.search(r"(?:current\s+|present\s+)?(?:location|city|based\s+(?:in|at))\s*[:\-]?\s+([A-Za-z][A-Za-z .,]{1,60})", line, re.I)
        if match and not re.search(r"preferred", line, re.I):
            return match.group(1).strip(" .,")
    for line in lines[:10]:
        for city in CITIES:
            if re.search(rf"\b{re.escape(city)}\b", line, re.I):
                return city
    return None


def _qualification(lines: list[str]) -> str | None:
    for label, pattern in DEGREES:
        for line in lines:
            match = pattern.search(line)
            if match:
                after = re.split(r"[,;]|\b(?:19|20)\d{2}\b", line[match.end():match.end() + 70])[0]
                spec = SPECIALIZATION.search(after)
                return f"{label} ({spec.group(1).title().replace(' And ', ' and ')})" if spec else label
    return None


def _split_skills(text: str) -> list[str]:
    items = re.split(r"\s*[,;|•·\t]\s*|\s{2,}", text)   # not "/": keeps PL/SQL, S/4HANA, CI/CD
    return [item.strip(" .") for item in items if 1 <= len(item.strip(" .")) <= 40]


def _skills(lines: list[str]) -> str | None:
    found: list[str] = []
    for index, line in enumerate(lines):
        match = SKILL_HEADING.match(line)
        if not match:
            continue
        if match.group(1):
            found.extend(_split_skills(match.group(1)))
        for following in lines[index + 1:index + 15]:
            if SECTION_HEADINGS.match(following) and not SKILL_HEADING.match(following):
                break
            value = following.split(":", 1)[1] if ":" in following and len(following.split(":", 1)[0]) <= 30 else following
            found.extend(_split_skills(value))
        if found:
            break
    if not found:
        text = "\n".join(lines)
        found = [skill for skill in KNOWN_SKILLS if re.search(rf"(?<![\w.]){re.escape(skill)}(?![\w])", text, re.I)]
    found = [skill for skill in found
             if not DATE_TOKEN.search(skill) and not CURRENT_MARK.search(skill) and not COMPANY_HINT.search(skill)]
    unique: list[str] = []
    for skill in found:
        if skill.lower() not in {u.lower() for u in unique} and not re.fullmatch(r"[\d\s.]+", skill):
            unique.append(skill)
    return ", ".join(unique[:25]) or None


FIELDS = ("full_name", "email", "phone", "total_experience", "relevant_experience", "current_employer",
          "current_location", "highest_qualification", "primary_skills")


def parse_lines(lines: list[str]) -> dict:
    total, relevant = _experience(lines)
    return {
        "full_name": _name(lines),
        "email": _email(lines),
        "phone": _phone(lines),
        "total_experience": total,
        "relevant_experience": relevant,
        "current_employer": _employer(lines),
        "current_location": _location(lines),
        "highest_qualification": _qualification(lines),
        "primary_skills": _skills(lines),
    }


def parse_resume(filename: str, content: bytes) -> dict:
    """Best values found across the document's text readings."""
    result = {field: None for field in FIELDS}
    for lines in extract_text_variants(filename, content):
        parsed = parse_lines(lines)
        for field in FIELDS:
            if result[field] is None and parsed[field] is not None:
                result[field] = parsed[field]
    return result
