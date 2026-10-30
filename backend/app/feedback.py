"""Bug reports and feature ideas sent from the app.

Anyone signed in can send feedback and follow the status of their own reports.
The Super Admin (users.manage) reviews everything, sets a status and can download it.
"""
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.db_dependencies import get_db
from app.dependencies import get_current_user, require_permission
from app.models import Feedback, User

router = APIRouter(prefix="/feedback", tags=["Feedback"])

KIND_LABEL = {"bug": "Bug", "feature": "Feature idea", "other": "Other"}
STATUS_LABEL = {"new": "New", "in_progress": "In progress", "done": "Done", "wont_fix": "Won't fix"}
ROLE_LABEL = {"Super Admin": "Super Admin", "Recruiter": "Recruiter", "Vendor": "Channel partner", "Client": "Client"}


class FeedbackCreate(BaseModel):
    kind: Literal["bug", "feature", "other"]
    title: str = Field(min_length=3, max_length=200)
    details: str = Field(min_length=1, max_length=5000)
    page: str | None = Field(default=None, max_length=200)
    browser: str | None = Field(default=None, max_length=300)


class FeedbackUpdate(BaseModel):
    status: Literal["new", "in_progress", "done", "wont_fix"] | None = None
    admin_note: str | None = Field(default=None, max_length=5000)


def _out(item: Feedback) -> dict:
    user = item.user
    return {
        "id": item.id,
        "kind": item.kind,
        "title": item.title,
        "details": item.details,
        "page": item.page,
        "browser": item.browser,
        "status": item.status,
        "admin_note": item.admin_note,
        "created_at": item.created_at,
        "updated_at": item.updated_at,
        "user_name": (user.full_name or user.email) if user else None,
        "user_email": user.email if user else None,
        "user_role": ", ".join(ROLE_LABEL.get(role.name, role.name) for role in user.roles) if user else None,
        "user_company": (user.vendor.name if user.vendor else user.client.company_name if user.client else None) if user else None,
    }


@router.post("", status_code=201)
def send_feedback(payload: FeedbackCreate, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    title, details = payload.title.strip(), payload.details.strip()
    if len(title) < 3 or not details:
        raise HTTPException(status_code=422, detail="Please give a short title and describe what happened or what you'd like.")
    item = Feedback(
        organization_id=current_user.organization_id,
        user_id=current_user.id,
        kind=payload.kind,
        title=title,
        details=details,
        page=(payload.page or "").strip() or None,
        browser=(payload.browser or "").strip() or None,
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return _out(item)


@router.get("/mine")
def my_feedback(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    items = db.scalars(
        select(Feedback).options(selectinload(Feedback.user)).where(Feedback.user_id == current_user.id).order_by(Feedback.id.desc())
    ).all()
    return [_out(item) for item in items]


def _query(current_user: User, status: str | None, kind: str | None):
    query = (
        select(Feedback)
        .options(selectinload(Feedback.user).selectinload(User.roles))
        .where(Feedback.organization_id == current_user.organization_id)
        .order_by(Feedback.id.desc())
    )
    if status:
        query = query.where(Feedback.status == status)
    if kind:
        query = query.where(Feedback.kind == kind)
    return query


@router.get("")
def list_feedback(
    status: Literal["new", "in_progress", "done", "wont_fix"] | None = None,
    kind: Literal["bug", "feature", "other"] | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("users.manage")),
):
    return [_out(item) for item in db.scalars(_query(current_user, status, kind)).all()]


@router.patch("/{feedback_id}")
def update_feedback(
    feedback_id: int,
    payload: FeedbackUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("users.manage")),
):
    item = db.get(Feedback, feedback_id)
    if item is None or item.organization_id != current_user.organization_id:
        raise HTTPException(status_code=404, detail="Feedback not found")
    changes = payload.model_dump(exclude_unset=True)
    if "status" in changes and changes["status"] is not None:
        item.status = changes["status"]
    if "admin_note" in changes:
        item.admin_note = (changes["admin_note"] or "").strip() or None
    item.updated_at = datetime.now(timezone.utc)
    db.commit()
    db.refresh(item)
    return _out(item)


@router.get("/export")
def export_feedback(
    status: Literal["new", "in_progress", "done", "wont_fix"] | None = None,
    kind: Literal["bug", "feature", "other"] | None = None,
    tz_offset: int = Query(default=0, ge=-840, le=840),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("users.manage")),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    items = [_out(item) for item in db.scalars(_query(current_user, status, kind)).all()]
    shift = timedelta(minutes=tz_offset)
    wb = Workbook()
    ws = wb.active
    ws.title = "Feedback"
    ws["A1"] = "Feedback (bugs and feature ideas)"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="173153")
    now = datetime.now(timezone.utc) + shift
    ws["A2"] = f"Generated {now:%d %b %Y, %I:%M %p}   ·   Status: {STATUS_LABEL.get(status, 'All')}   ·   Type: {KIND_LABEL.get(kind, 'All')}   ·   {len(items)} item(s)"
    ws["A2"].font = Font(name="Arial", size=9, color="5B6B80")
    headers = ["S.No", "ID", "Type", "Status", "Title", "Details", "From", "Role", "Company", "Page", "Browser", "Sent on", "Admin note"]
    for index, label in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=index, value=label)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173153")
    body = Font(name="Arial", size=10)
    for offset, item in enumerate(items):
        row = 5 + offset
        values = [
            offset + 1, item["id"], KIND_LABEL[item["kind"]], STATUS_LABEL[item["status"]], item["title"], item["details"],
            item["user_name"], item["user_role"], item["user_company"], item["page"], item["browser"],
            (item["created_at"] + shift).replace(tzinfo=None) if item["created_at"] else None, item["admin_note"],
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=index, value=value)
            cell.font = body
            cell.alignment = Alignment(vertical="top", wrap_text=index in (5, 6, 13), horizontal="center" if index == 1 else None)
            if index == 12:
                cell.number_format = "dd mmm yyyy hh:mm"
    widths = [7, 7, 13, 12, 34, 60, 22, 16, 18, 16, 30, 18, 40]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A5"
    if items:
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{4 + len(items)}"
    buffer = BytesIO()
    wb.save(buffer)
    return StreamingResponse(
        BytesIO(buffer.getvalue()),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="feedback-{now:%Y-%m-%d}.xlsx"'},
    )
