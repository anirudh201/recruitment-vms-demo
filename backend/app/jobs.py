import os
import re
import uuid
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from fastapi.responses import FileResponse
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.db_dependencies import get_db
from app.dependencies import get_current_user, require_permission
from app.models import Client, Job, JobVendorAssignment, User, Vendor
from app.partner_access import hide_clients_from, is_client_user, is_partner_user, permission_codes


router = APIRouter(prefix="/jobs", tags=["Job Management"])

# Original JD documents are stored outside the database. Set JD_UPLOAD_DIR in
# backend/.env to move them (for example to a mounted volume in production).
JD_UPLOAD_DIR = Path(
    os.getenv("JD_UPLOAD_DIR")
    or Path(__file__).resolve().parent.parent / "uploads" / "jd"
)
MAX_JD_BYTES = 10 * 1024 * 1024
JD_FILE_TYPES = {
    ".pdf": ("application/pdf", b"%PDF-"),
    ".docx": (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        b"PK\x03\x04",
    ),
}


def _extract_pdf_page_text(page) -> str:
    """Layout mode keeps words on their visual line; fall back to the default mode."""
    for kwargs in ({"extraction_mode": "layout"}, {}):
        try:
            text = page.extract_text(**kwargs) or ""
        except Exception:
            continue
        if text.strip():
            return text
    return ""


def _reflow_fragmented_lines(lines: list[str]) -> list[str]:
    """Some designed PDFs extract as one word per line. Rebuild sentences in that case only."""
    if len(lines) < 20:
        return lines
    short = sum(1 for line in lines if len(line.split()) <= 2)
    if short / len(lines) < 0.6:
        return lines

    merged: list[str] = []
    buffer: list[str] = []
    for line in lines:
        buffer.append(line)
        if re.search(r"[.!?]$", line) and len(buffer) >= 4:
            merged.append(" ".join(buffer))
            buffer = []
    if buffer:
        merged.append(" ".join(buffer))
    return merged


def _jd_file_path(stored_name: str | None) -> Path | None:
    """Resolve a stored JD file name, refusing anything outside JD_UPLOAD_DIR."""
    if not stored_name:
        return None
    base = JD_UPLOAD_DIR.resolve()
    candidate = (base / stored_name).resolve()
    return candidate if base in candidate.parents else None


def _delete_jd_file(stored_name: str | None) -> None:
    path = _jd_file_path(stored_name)
    if path is not None:
        path.unlink(missing_ok=True)


@router.post("/parse-jd")
async def parse_job_description(
    file: UploadFile = File(...),
    current_user: User = Depends(require_permission("jobs.manage")),
):
    """Extract text from a PDF/DOCX job description and infer only explicitly stated fields."""
    filename = (file.filename or "").lower()
    if not filename.endswith((".pdf", ".docx")):
        raise HTTPException(status_code=415, detail="Upload a PDF or DOCX job description.")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="The file must be 10 MB or smaller.")

    try:
        from io import BytesIO
        if filename.endswith(".pdf"):
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            text = "\n".join(_extract_pdf_page_text(page) for page in reader.pages)
        else:
            from docx import Document
            document = Document(BytesIO(content))
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
            text = "\n".join(paragraphs)
    except Exception as exc:
        raise HTTPException(status_code=422, detail="Unable to read this file. Please use a text-based PDF or DOCX.") from exc

    text = text.replace("\r", "\n")
    # PDFs often emit bullets as control / private-use characters (e.g. \x7f, U+F0B7).
    text = re.sub(r"[\x0b\x0c]", "\n", text)
    text = re.sub(r"[\x00-\x08\x0e-\x1f\x7f\uf000-\uf8ff]", "", text)
    lines = [re.sub(r"\s+", " ", line).strip(" \t•*-:") for line in text.splitlines()]
    lines = [line for line in lines if line]
    lines = _reflow_fragmented_lines(lines)
    normalized = "\n".join(lines)
    if not normalized.strip():
        raise HTTPException(status_code=422, detail="No readable text was found. This may be a scanned PDF.")

    def label_value(labels: str) -> str | None:
        pattern = re.compile(rf"(?im)^\s*(?:{labels})\s*[:\-]\s*(.+?)\s*$")
        match = pattern.search(normalized)
        return match.group(1).strip() if match else None

    def section_value(labels: str, stop_labels: str) -> str | None:
        pattern = re.compile(rf"(?ims)^\s*(?:{labels})\s*[:\-]?\s*\n?(.*?)(?=^\s*(?:{stop_labels})\s*[:\-]?\s*$|\Z)")
        match = pattern.search(normalized)
        value = match.group(1).strip() if match else ""
        return value[:4000] or None

    title = label_value(r"job\s*title|position(?:\s*title)?|role")
    if not title:
        for line in lines[:12]:
            if re.search(r"\b(job description|job title|requirement|experience|qualification|responsibilities)\b", line, re.I):
                continue
            if 3 <= len(line) <= 120:
                title = line
                break

    exp = re.search(r"(?i)(\d{1,2})\s*(?:-|to|–|—)\s*(\d{1,2})\s*\+?\s*(?:years?|yrs?)(?:\s+of\s+experience)?", normalized)
    if not exp:
        exp = re.search(r"(?i)(\d{1,2})\s*\+\s*(?:years?|yrs?)(?:\s+of\s+experience)?", normalized)
    min_exp = max_exp = None
    if exp:
        min_exp = int(exp.group(1))
        max_exp = int(exp.group(2)) if exp.lastindex and exp.lastindex >= 2 else None

    location = label_value(r"work\s*location|job\s*location|location|based\s*at")
    qualification = label_value(r"qualification|education|educational\s*qualification|degree")
    arrangement_raw = label_value(r"work\s*arrangement|work\s*mode|work\s*type") or ""
    if not arrangement_raw:
        arrangement_match = re.search(r"(?i)\b(on[- ]site|hybrid|remote)\b", normalized)
        arrangement_raw = arrangement_match.group(1) if arrangement_match else ""
    arrangement = None
    if arrangement_raw:
        low = arrangement_raw.lower()
        arrangement = "On-site" if "site" in low else ("Hybrid" if "hybrid" in low else ("Remote" if "remote" in low else None))

    openings_raw = label_value(r"no\.?\s*of\s*(?:openings|positions)|number\s*of\s*(?:openings|positions)|openings|positions|vacancies")
    openings_match = re.search(r"\d+", openings_raw or "")
    openings = int(openings_match.group()) if openings_match else None

    mandatory = section_value(r"mandatory\s*skills|required\s*skills|must[- ]have\s*skills|technical\s*skills", r"preferred\s*skills|nice[- ]to[- ]have|responsibilities|roles?\s*&?\s*responsibilities|job\s*description|qualification|education|location|experience|benefits|about\s*the\s*role")
    preferred = section_value(r"preferred\s*skills|nice[- ]to[- ]have", r"responsibilities|roles?\s*&?\s*responsibilities|job\s*description|qualification|education|location|experience|benefits|about\s*the\s*role|mandatory\s*skills|required\s*skills")

    return {
        "title": title,
        "description": normalized[:12000],
        "mandatory_skills": mandatory,
        "preferred_skills": preferred,
        "min_experience": min_exp,
        "max_experience": max_exp,
        "qualification": qualification,
        "location": location,
        "work_arrangement": arrangement,
        "budget_min": None,
        "budget_max": None,
        "openings": openings,
    }


class JobCreate(BaseModel):
    title: str = Field(min_length=1, max_length=200)
    client_id: int | None = Field(default=None, gt=0)
    description: str = Field(min_length=1)
    mandatory_skills: str | None = None
    preferred_skills: str | None = None
    min_experience: int = Field(default=0, ge=0)
    max_experience: int | None = Field(default=None, ge=0)
    qualification: str | None = Field(default=None, max_length=300)
    location: str = Field(min_length=1, max_length=200)
    work_arrangement: str | None = Field(default=None, max_length=50)
    budget_min: int | None = Field(default=None, ge=0)
    budget_max: int | None = Field(default=None, ge=0)
    openings: int = Field(default=1, gt=0)
    status: str = Field(default="draft", pattern="^(draft|open|closed)$")
    recruiter_id: int | None = Field(default=None, gt=0)

    def validate_ranges(self):
        if (
            self.max_experience is not None
            and self.max_experience < self.min_experience
        ):
            raise ValueError("max_experience cannot be less than min_experience")
        if (
            self.budget_min is not None
            and self.budget_max is not None
            and self.budget_max < self.budget_min
        ):
            raise ValueError("budget_max cannot be less than budget_min")


class JobUpdate(BaseModel):
    title: str | None = Field(default=None, min_length=1, max_length=200)
    client_id: int | None = Field(default=None, gt=0)
    description: str | None = Field(default=None, min_length=1)
    mandatory_skills: str | None = None
    preferred_skills: str | None = None
    min_experience: int | None = Field(default=None, ge=0)
    max_experience: int | None = Field(default=None, ge=0)
    qualification: str | None = Field(default=None, max_length=300)
    location: str | None = Field(default=None, min_length=1, max_length=200)
    work_arrangement: str | None = Field(default=None, max_length=50)
    budget_min: int | None = Field(default=None, ge=0)
    budget_max: int | None = Field(default=None, ge=0)
    openings: int | None = Field(default=None, gt=0)
    status: str | None = Field(default=None, pattern="^(draft|open|closed)$")
    is_active: bool | None = None
    recruiter_id: int | None = Field(default=None, gt=0)


class JobResponse(BaseModel):
    model_config = ConfigDict(from_attributes=True)

    id: int
    organization_id: int
    created_by_user_id: int
    client_id: int | None
    title: str
    description: str
    mandatory_skills: str | None
    preferred_skills: str | None
    min_experience: int
    max_experience: int | None
    qualification: str | None
    location: str
    work_arrangement: str | None
    budget_min: int | None
    budget_max: int | None
    openings: int
    status: str
    created_at: datetime
    updated_at: datetime
    jd_file_name: str | None = None
    client_name: str | None = None
    is_active: bool = True
    recruiter_id: int | None = None
    recruiter_name: str | None = None


def is_super_admin(user: User) -> bool:
    return any(role.name == "Super Admin" for role in user.roles)


def can_view_job(db: Session, user: User, job: Job) -> bool:
    if is_super_admin(user):
        return True

    if user.organization_id != job.organization_id:
        return False

    if not job.is_active and "jobs.manage" not in permission_codes(user):
        return False

    if is_client_user(user):
        # A client company only sees its own requirements, and not drafts still being prepared.
        return user.client_id is not None and job.client_id == user.client_id and job.status != "draft"

    if any(role.name == "Vendor" for role in user.roles):
        if user.vendor_id is None:
            return False

        assignment = db.scalar(
            select(JobVendorAssignment.id).where(
                JobVendorAssignment.job_id == job.id,
                JobVendorAssignment.vendor_id == user.vendor_id,
                JobVendorAssignment.is_active.is_(True),
            )
        )
        return assignment is not None

    return any(role.name == "Recruiter" for role in user.roles)


@router.post(
    "",
    response_model=JobResponse,
    status_code=status.HTTP_201_CREATED,
)
def create_job(
    payload: JobCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("jobs.manage")),
):
    if payload.max_experience is not None and (
        payload.max_experience < payload.min_experience
    ):
        raise HTTPException(
            status_code=422,
            detail="max_experience cannot be less than min_experience",
        )

    if (
        payload.budget_min is not None
        and payload.budget_max is not None
        and payload.budget_max < payload.budget_min
    ):
        raise HTTPException(
            status_code=422,
            detail="budget_max cannot be less than budget_min",
        )
    if payload.client_id is not None:
        client = db.get(Client, payload.client_id)

        if (
            client is None
            or client.organization_id != current_user.organization_id
        ):
            raise HTTPException(
                status_code=422,
                detail="Client not found in your organization",
            )
        if not client.is_active:
            raise HTTPException(status_code=422, detail="This client is deactivated. Activate it first.")
    _check_recruiter(db, current_user, payload.recruiter_id)
    job_data = payload.model_dump()
    if not is_super_admin(current_user):
        # Non-Super Admin users may create a job only as a draft.
        job_data["status"] = "draft"

    job = Job(
        organization_id=current_user.organization_id,
        created_by_user_id=current_user.id,
        **job_data,
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    return _job_for(current_user, job)


@router.get("", response_model=list[JobResponse])
def list_jobs(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("jobs.view")),
):
    jobs = db.scalars(select(Job).order_by(Job.id.desc())).all()
    return [_job_for(current_user, job) for job in jobs if can_view_job(db, current_user, job)]


def _job_for(user: User, job: Job) -> JobResponse:
    """Job as this user may see it: partners do not get the client link unless allowed."""
    response = JobResponse.model_validate(job)
    recruiter = job.recruiter
    staff_view = not is_partner_user(user) and not is_client_user(user)
    recruiter_fields = {
        "recruiter_id": job.recruiter_id if staff_view else None,
        "recruiter_name": (recruiter.full_name or recruiter.email) if staff_view and recruiter is not None else None,
    }
    if hide_clients_from(user):
        return response.model_copy(update={"client_id": None, "client_name": None, **recruiter_fields})
    return response.model_copy(update={"client_name": job.client.company_name if job.client is not None else None, **recruiter_fields})


def _check_recruiter(db: Session, current_user: User, recruiter_id: int | None) -> None:
    """Only an active Recruiter of the same organization can own a job."""
    if recruiter_id is None:
        return
    recruiter = db.get(User, recruiter_id)
    if (
        recruiter is None
        or recruiter.organization_id != current_user.organization_id
        or not recruiter.is_active
        or not any(role.name == "Recruiter" for role in recruiter.roles)
    ):
        raise HTTPException(status_code=422, detail="Choose an active recruiter from your organization.")


class JobAssignmentResponse(BaseModel):
    job_id: int
    vendor_id: int
    vendor_name: str
    assigned_at: datetime


@router.get("/assignments", response_model=list[JobAssignmentResponse])
def list_job_assignments(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("jobs.manage")),
):
    """Every active partner assignment in the organization (job <-> vendor)."""
    rows = db.execute(
        select(JobVendorAssignment, Vendor.name)
        .join(Job, Job.id == JobVendorAssignment.job_id)
        .join(Vendor, Vendor.id == JobVendorAssignment.vendor_id)
        .where(
            Job.organization_id == current_user.organization_id,
            JobVendorAssignment.is_active.is_(True),
        )
        .order_by(JobVendorAssignment.assigned_at.desc())
    ).all()
    return [
        JobAssignmentResponse(
            job_id=assignment.job_id,
            vendor_id=assignment.vendor_id,
            vendor_name=vendor_name,
            assigned_at=assignment.assigned_at,
        )
        for assignment, vendor_name in rows
    ]


@router.get("/permissions")
def get_job_permissions(
    current_user: User = Depends(require_permission("jobs.view")),
):
    can_manage_jobs = any(
        permission.code == "jobs.manage"
        for role in current_user.roles
        for permission in role.permissions
    )
    return {
        "can_change_status": is_super_admin(current_user),
        "can_manage_jobs": can_manage_jobs,
    }


@router.get("/{job_id}", response_model=JobResponse)
def get_job(
    job_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("jobs.view")),
):
    job = db.get(Job, job_id)

    if job is None or not can_view_job(db, current_user, job):
        raise HTTPException(status_code=404, detail="Job not found")

    return _job_for(current_user, job)


@router.patch("/{job_id}", response_model=JobResponse)
def update_job(
    job_id: int,
    payload: JobUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("jobs.manage")),
):
    job = db.get(Job, job_id)

    if job is None:
        raise HTTPException(status_code=404, detail="Job not found")

    if job.organization_id != current_user.organization_id:
        raise HTTPException(status_code=404, detail="Job not found")

    changes = payload.model_dump(exclude_unset=True)

    if "status" in changes and not is_super_admin(current_user):
        raise HTTPException(
            status_code=403,
            detail="Only Super Admin can change job status",
        )

    if "client_id" in changes and changes["client_id"] is not None:
        client = db.get(Client, changes["client_id"])

        if (
            client is None
            or client.organization_id != current_user.organization_id
        ):
            raise HTTPException(
                status_code=422,
                detail="Client not found in your organization",
            )
        if not client.is_active and changes["client_id"] != job.client_id:
            raise HTTPException(status_code=422, detail="This client is deactivated. Activate it first.")

    if changes.get("recruiter_id") is not None and changes["recruiter_id"] != job.recruiter_id:
        _check_recruiter(db, current_user, changes["recruiter_id"])

    min_experience = changes.get("min_experience", job.min_experience)
    max_experience = changes.get("max_experience", job.max_experience)
    budget_min = changes.get("budget_min", job.budget_min)
    budget_max = changes.get("budget_max", job.budget_max)

    if max_experience is not None and max_experience < min_experience:
        raise HTTPException(status_code=422, detail="Invalid experience range")

    if budget_min is not None and budget_max is not None and budget_max < budget_min:
        raise HTTPException(status_code=422, detail="Invalid budget range")

    for field, value in changes.items():
        setattr(job, field, value)

    db.commit()
    db.refresh(job)
    return _job_for(current_user, job)


@router.post("/{job_id}/jd-file", response_model=JobResponse)
async def upload_job_jd_file(
    job_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("jobs.manage")),
):
    """Attach (or replace) the original PDF/DOCX job description on a job."""
    job = db.get(Job, job_id)
    if job is None or job.organization_id != current_user.organization_id:
        raise HTTPException(status_code=404, detail="Job not found")

    original_name = Path(file.filename or "").name
    extension = Path(original_name).suffix.lower()
    if extension not in JD_FILE_TYPES:
        raise HTTPException(status_code=415, detail="Upload a PDF or DOCX job description.")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > MAX_JD_BYTES:
        raise HTTPException(status_code=413, detail="The file must be 10 MB or smaller.")

    content_type, signature = JD_FILE_TYPES[extension]
    if not content.startswith(signature):
        raise HTTPException(
            status_code=422,
            detail="The file content does not match its extension. Please upload a real PDF or DOCX.",
        )

    JD_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid.uuid4().hex}{extension}"
    (JD_UPLOAD_DIR / stored_name).write_bytes(content)

    previous_file = job.jd_file_path
    job.jd_file_name = (original_name or f"JD{extension}")[:300]
    job.jd_file_path = stored_name
    job.jd_content_type = content_type

    try:
        db.commit()
    except Exception:
        db.rollback()
        _delete_jd_file(stored_name)
        raise

    db.refresh(job)
    if previous_file:
        _delete_jd_file(previous_file)
    return _job_for(current_user, job)


@router.get("/{job_id}/jd-file")
def download_job_jd_file(
    job_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("jobs.view")),
):
    """Return the original uploaded JD document for a job the user may view."""
    job = db.get(Job, job_id)
    if job is None or not can_view_job(db, current_user, job):
        raise HTTPException(status_code=404, detail="Job not found")

    if not job.jd_file_path:
        raise HTTPException(status_code=404, detail="No JD file has been uploaded for this job.")

    path = _jd_file_path(job.jd_file_path)
    if path is None or not path.is_file():
        raise HTTPException(status_code=404, detail="The stored JD file could not be found on the server.")

    return FileResponse(
        path,
        media_type=job.jd_content_type or "application/octet-stream",
        filename=job.jd_file_name or path.name,
        content_disposition_type="inline",
        headers={"X-Content-Type-Options": "nosniff"},
    )


@router.post("/{job_id}/assign-vendor/{vendor_id}", status_code=201)
def assign_vendor(
    job_id: int,
    vendor_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("jobs.manage")),
):
    job = db.get(Job, job_id)
    vendor = db.get(Vendor, vendor_id)

    if job is None or job.organization_id != current_user.organization_id:
        raise HTTPException(status_code=404, detail="Job not found")
    if not job.is_active:
        raise HTTPException(status_code=409, detail="This job is deactivated. Activate it before assigning partners.")

    if (
        vendor is None
        or vendor.organization_id != current_user.organization_id
        or not vendor.is_active
    ):
        raise HTTPException(status_code=404, detail="Active vendor not found")

    existing = db.scalar(
        select(JobVendorAssignment).where(
            JobVendorAssignment.job_id == job_id,
            JobVendorAssignment.vendor_id == vendor_id,
        )
    )

    if existing is not None:
        if existing.is_active:
            raise HTTPException(status_code=409, detail="Vendor already assigned")
        existing.is_active = True
        existing.assigned_at = datetime.utcnow()
    else:
        db.add(JobVendorAssignment(job_id=job_id, vendor_id=vendor_id))

    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(status_code=409, detail="Could not assign vendor")

    return {"message": "Vendor assigned successfully", "job_id": job_id, "vendor_id": vendor_id}


@router.delete("/{job_id}/assign-vendor/{vendor_id}")
def unassign_vendor(
    job_id: int,
    vendor_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_permission("jobs.manage")),
):
    """Stop a partner working on a job. The assignment is deactivated, not deleted,
    so earlier submissions and their history stay intact."""
    job = db.get(Job, job_id)
    if job is None or job.organization_id != current_user.organization_id:
        raise HTTPException(status_code=404, detail="Job not found")

    assignment = db.scalar(
        select(JobVendorAssignment).where(
            JobVendorAssignment.job_id == job_id,
            JobVendorAssignment.vendor_id == vendor_id,
            JobVendorAssignment.is_active.is_(True),
        )
    )
    if assignment is None:
        raise HTTPException(status_code=404, detail="This partner is not assigned to the job")

    assignment.is_active = False
    db.commit()
    return {"message": "Vendor unassigned successfully", "job_id": job_id, "vendor_id": vendor_id}
